"""Ingest and clean a tabular spend file (CSV / XLSX) into the DuckDB star schema.

Designed for messy, "not properly arranged" inputs:
- preamble/title rows before the real header (auto-detected),
- odd column names (alias matching, with an optional SLM fallback),
- currency symbols and thousands separators in amounts,
- mixed date formats,
- blank rows, and duplicate supplier/category/BU names (de-duplicated).
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import os
import re

import duckdb

from config import settings
from db import DB_LOCK

HERE = os.path.dirname(os.path.abspath(__file__))
SCHEMA_PATH = os.path.join(HERE, "scripts", "sql", "schema.sql")

FIELD_ALIASES = {
    "supplier_name": ["supplier_name", "supplier", "vendor", "vendor_name", "supplier name", "vendor name", "payee"],
    "category": ["category", "category_level2", "commodity", "spend_category", "category name", "gl category"],
    "category_level1": ["category_level1", "category_l1", "parent_category"],
    "category_code": ["category_code", "unspsc", "code"],
    "business_unit": ["business_unit", "bu", "bu_name", "department", "cost_center", "business unit", "cost centre"],
    "region": ["region", "geo"],
    "invoice_date": ["invoice_date", "date", "posting_date", "document_date", "invoice date", "posting date"],
    "qty": ["qty", "quantity", "units"],
    "unit_price": ["unit_price", "price", "unit price", "rate"],
    "amount": ["amount", "spend", "total", "net_amount", "value", "amount_usd", "net amount", "gross amount"],
    "item": ["item", "item name", "item description", "description", "product", "product name", "line item", "sku"],
    "currency": ["currency", "curr"],
    "doc_type": ["doc_type", "document_type", "type"],
    "doc_number": ["doc_number", "document_number", "invoice_number", "po_number", "doc no", "invoice no"],
    "country": ["country", "supplier_country"],
    "is_preferred": ["is_preferred", "preferred", "preferred_supplier"],
    "contract_id": ["contract_id", "contract", "contract_no", "contract number"],
    "contracted_unit_price": ["contracted_unit_price", "contract_price", "contracted price"],
    "is_maverick": ["is_maverick", "maverick", "off_contract"],
}
ALL_ALIASES = {a for aliases in FIELD_ALIASES.values() for a in aliases}
DATE_FORMATS = ["%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d", "%d-%b-%Y", "%m-%d-%Y", "%d.%m.%Y", "%b %d, %Y"]


def _norm(s) -> str:
    return str(s).strip().lower()


def _canon(s) -> str:
    """Separator-insensitive key: lower-cased, with runs of spaces / underscores /
    dashes collapsed to a single space (so 'Spend Category' == 'spend_category')."""
    return re.sub(r"[\s_\-]+", " ", str(s).strip().lower())


ALL_ALIASES_CANON = {_canon(a) for a in ALL_ALIASES}


def _num(s):
    if s is None:
        return None
    cleaned = re.sub(r"[^0-9.\-]", "", str(s))
    if cleaned in ("", "-", "."):
        return None
    try:
        return float(cleaned)
    except ValueError:
        return None


def _date(s):
    s = (str(s) if s is not None else "").strip()
    if not s:
        return None
    try:
        return dt.date.fromisoformat(s[:10])
    except ValueError:
        pass
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _bool(s) -> bool:
    return str(s).strip().lower() in {"1", "true", "yes", "y", "t"}


def read_tabular(data: bytes, filename: str) -> list[list[str]]:
    """Return the file as a list of rows (each a list of cell strings)."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    # Sniff the delimiter so comma, semicolon, tab (TSV) and pipe files all parse.
    first_line = next((ln for ln in text.splitlines() if ln.strip()), "")
    counts = {d: first_line.count(d) for d in (",", ";", "\t", "|")}
    delimiter = max(counts, key=counts.get)
    if counts[delimiter] == 0:
        delimiter = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def detect_header_row(rows: list[list[str]]) -> int:
    """Pick the row (within the first 15) that best looks like a header."""
    best_i, best_score = 0, -1
    for i, row in enumerate(rows[:15]):
        cells = [_canon(c) for c in row]
        hits = sum(1 for c in cells if c in ALL_ALIASES_CANON)
        nonempty = sum(1 for c in cells if c)
        score = hits * 100 + nonempty
        if hits > 0 and score > best_score:
            best_i, best_score = i, score
    return best_i


def map_columns(header: list[str]) -> dict[str, int]:
    lower = {_canon(h): idx for idx, h in enumerate(header)}
    mapping: dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if _canon(alias) in lower:
                mapping[field] = lower[_canon(alias)]
                break
    return mapping


def infer_mapping_with_slm(header: list[str], sample_rows: list[list[str]]) -> dict[str, int]:
    """Ask the SLM to map unknown columns to canonical fields (best effort)."""
    from llm import chat

    prompt = (
        "Map spreadsheet columns to canonical procurement fields.\n"
        f"Columns (index: name): {list(enumerate(header))}\n"
        f"Sample rows: {sample_rows[:3]}\n"
        "Canonical fields: supplier_name, category, business_unit, invoice_date, qty, "
        "unit_price, amount, currency, doc_type, doc_number, contract_id.\n"
        'Return ONLY JSON mapping canonical field -> column index for columns that clearly match, '
        'e.g. {"supplier_name": 0, "amount": 3}.'
    )
    raw = chat([{"role": "user", "content": prompt}])
    raw = re.sub(r"^```[a-zA-Z]*|```$", "", raw.strip()).strip()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return {}
    lower_headers = {_norm(h): i for i, h in enumerate(header)}
    out: dict[str, int] = {}
    for key, val in data.items():
        if isinstance(val, bool):
            continue
        if isinstance(val, (int, float)) and 0 <= int(val) < len(header):
            out[key] = int(val)
        elif isinstance(val, str):
            vs = val.strip()
            if vs.isdigit() and 0 <= int(vs) < len(header):
                out[key] = int(vs)
            elif _norm(vs) in lower_headers:
                out[key] = lower_headers[_norm(vs)]
    return out


def ensure_schema(con) -> None:
    with open(SCHEMA_PATH, encoding="utf-8") as f:
        sql = re.sub(r"--[^\n]*", "", f.read())
    for stmt in (s.strip() for s in sql.split(";") if s.strip()):
        con.execute(stmt)


def _cell(row: list[str], idx):
    return row[idx] if idx is not None and idx < len(row) else None


def _load(con, mapping, header, data_rows, mode, notes) -> dict:
    if mode == "append":
        suppliers = {n.casefold(): k for k, n in con.execute(
            "SELECT supplier_key, supplier_name FROM dim_supplier").fetchall()}
        categories = {n.casefold(): k for k, n in con.execute(
            "SELECT category_key, level2 FROM dim_category").fetchall()}
        bus = {n.casefold(): k for k, n in con.execute(
            "SELECT bu_key, bu_name FROM dim_business_unit").fetchall()}
        contracts = {c: k for k, c in con.execute(
            "SELECT contract_key, contract_id FROM dim_contract").fetchall()}

        def nx(table, col):
            return con.execute(f"SELECT COALESCE(MAX({col}), 0) FROM {table}").fetchone()[0] + 1

        sk, ck, bk, ctk, fk = (nx("dim_supplier", "supplier_key"), nx("dim_category", "category_key"),
                               nx("dim_business_unit", "bu_key"), nx("dim_contract", "contract_key"),
                               nx("fact_spend", "spend_id"))
    else:
        for table in ["fact_spend", "dim_contract", "dim_business_unit", "dim_category", "dim_supplier"]:
            con.execute(f"DELETE FROM {table}")
        suppliers, categories, bus, contracts = {}, {}, {}, {}
        sk = ck = bk = ctk = fk = 1

    new_s, new_c, new_b, new_ct, facts = [], [], [], [], []
    skipped = bad_dates = 0

    def g(row, field):
        return _cell(row, mapping.get(field))

    for row in data_rows:
        supplier = (g(row, "supplier_name") or "").strip()
        if not supplier:
            skipped += 1
            continue
        qty = _num(g(row, "qty")) or 1
        unit_price = _num(g(row, "unit_price"))
        amount = _num(g(row, "amount"))
        if amount is None and unit_price is not None:
            amount = round(qty * unit_price, 2)
        if amount is None:
            skipped += 1
            continue

        cat2 = (g(row, "category") or "Uncategorised").strip() or "Uncategorised"
        cat1 = (g(row, "category_level1") or cat2).strip() or cat2
        code = (g(row, "category_code") or "").strip() or None
        bu = (g(row, "business_unit") or "Unknown").strip() or "Unknown"
        region = (g(row, "region") or "").strip() or None
        currency = (g(row, "currency") or "USD").strip() or "USD"
        doc_type = (g(row, "doc_type") or "AP").strip() or "AP"
        doc_number = (g(row, "doc_number") or "").strip() or None
        item = (g(row, "item") or "").strip() or None
        country = (g(row, "country") or "").strip() or None
        is_pref = _bool(g(row, "is_preferred")) if "is_preferred" in mapping else False
        raw_date = g(row, "invoice_date")
        inv_date = _date(raw_date)
        if raw_date and str(raw_date).strip() and inv_date is None:
            bad_dates += 1
        contract_id = (g(row, "contract_id") or "").strip()
        contracted = _num(g(row, "contracted_unit_price"))

        key = supplier.casefold()
        if key not in suppliers:
            suppliers[key] = sk
            new_s.append((sk, supplier[:3].upper(), supplier, None, country, is_pref))
            sk += 1
        supplier_key = suppliers[key]

        ckey = cat2.casefold()
        if ckey not in categories:
            categories[ckey] = ck
            new_c.append((ck, code, cat1, cat2, cat2, cat2))
            ck += 1
        category_key = categories[ckey]

        bkey = bu.casefold()
        if bkey not in bus:
            bus[bkey] = bk
            new_b.append((bk, bu[:6].upper(), bu, region))
            bk += 1
        bu_key = bus[bkey]

        contract_key = None
        if contract_id:
            if contract_id not in contracts:
                contracts[contract_id] = ctk
                new_ct.append((ctk, contract_id, supplier_key, category_key, None, None, contracted, None, currency))
                ctk += 1
            contract_key = contracts[contract_id]

        is_mav = _bool(g(row, "is_maverick")) if "is_maverick" in mapping else (contract_key is None)
        facts.append((fk, doc_type, doc_number, 1, supplier_key, category_key, bu_key, contract_key,
                      inv_date, qty, unit_price, amount, currency, None, is_mav, "UPLOAD", item))
        fk += 1

    if new_s:
        con.executemany("INSERT INTO dim_supplier (supplier_key, supplier_id, supplier_name, parent_key, "
                        "country, is_preferred) VALUES (?,?,?,?,?,?)", new_s)
    if new_c:
        con.executemany("INSERT INTO dim_category (category_key, code, level1, level2, level3, level4) "
                        "VALUES (?,?,?,?,?,?)", new_c)
    if new_b:
        con.executemany("INSERT INTO dim_business_unit (bu_key, bu_code, bu_name, region) VALUES (?,?,?,?)", new_b)
    if new_ct:
        con.executemany("INSERT INTO dim_contract (contract_key, contract_id, supplier_key, category_key, "
                        "start_date, end_date, contracted_unit_price, committed_volume, currency) "
                        "VALUES (?,?,?,?,?,?,?,?,?)", new_ct)
    if facts:
        con.executemany("INSERT INTO fact_spend (spend_id, doc_type, doc_number, line_no, supplier_key, "
                        "category_key, bu_key, contract_key, invoice_date, qty, unit_price, amount, currency, "
                        "classification_confidence, is_maverick, source_system, item_name) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)", facts)

    total = con.execute("SELECT COUNT(*) FROM fact_spend").fetchone()[0]

    notes.append("Amounts cleaned of currency symbols and thousands separators.")
    notes.append("Supplier, category, and business-unit names de-duplicated (case-insensitive).")
    if bad_dates:
        notes.append(f"{bad_dates} row(s) had an unrecognised date and were kept with an empty date.")
    if skipped:
        notes.append(f"{skipped} row(s) skipped (missing a supplier or an amount).")

    return {
        "rows_read": len(data_rows),
        "rows_loaded": len(facts),
        "rows_skipped": skipped,
        "new_suppliers": len(new_s),
        "new_categories": len(new_c),
        "new_business_units": len(new_b),
        "new_contracts": len(new_ct),
        "total_rows_in_db": total,
    }


def ingest_bytes(data: bytes, filename: str, db_path: str | None = None,
                 mode: str = "replace", use_slm: bool = True) -> dict:
    """Read + clean + load a spend file. Returns a cleaning/loading summary."""
    db_path = db_path or settings.duckdb_path
    rows = [r for r in read_tabular(data, filename) if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("The file appears to be empty.")

    hdr_i = detect_header_row(rows)
    header = [str(c).strip() for c in rows[hdr_i]]
    data_rows = rows[hdr_i + 1:]

    mapping = map_columns(header)
    method = "column aliases"

    def has_required(m):
        return "supplier_name" in m and ("amount" in m or "unit_price" in m)

    if not has_required(mapping) and use_slm and settings.slm_api_key:
        slm_map = infer_mapping_with_slm(header, data_rows[:3])
        for k, v in slm_map.items():
            mapping.setdefault(k, v)
        if slm_map:
            method = "column aliases + SLM"

    if not has_required(mapping):
        raise ValueError(
            "Could not identify the required columns (a supplier column and an amount or unit-price "
            f"column). Detected headers: {header}"
        )

    notes: list[str] = []
    with DB_LOCK:
        con = duckdb.connect(db_path)
        try:
            ensure_schema(con)
            summary = _load(con, mapping, header, data_rows, mode, notes)
        finally:
            con.close()

    summary.update({
        "filename": os.path.basename(filename or "upload"),
        "mode": mode,
        "header_row_index": hdr_i,
        "mapping_method": method,
        "columns_mapped": {k: header[v] for k, v in sorted(mapping.items()) if v < len(header)},
        "notes": notes,
    })
    return summary
